"""КС-2 / КС-3 drafts: amounts from progress deltas between acts, VAT,
cumulative КС-3, period and deletion rules."""

import io

import pytest
from openpyxl import load_workbook

from scripts.make_demo_data import make_schedule_xml
from tests.test_issues import _admin


def test_ks_acts(client):
    h = _admin(client)
    pid = client.post("/api/projects", json={"name": "ЖК Демо"}, headers=h).json()["id"]
    assert client.post(f"/api/projects/{pid}/ks-acts", headers=h, json={"period_from": "2026-03-01", "period_to": "2026-03-31"}).status_code == 400
    tasks = client.post(f"/api/projects/{pid}/schedule/import", headers=h, files={"file": ("p.xml", make_schedule_xml(storeys=3))}).json()["tasks"]
    leaves = [t for t in tasks if not t["is_summary"]]
    expected = round(sum(t["cost"] * t["percent_complete"] / 100 for t in leaves if t["cost"]), 2)

    a1 = client.post(f"/api/projects/{pid}/ks-acts", headers=h, json={
        "period_from": "2026-03-01", "period_to": "2026-05-31", "customer": "ТОО «Заказчик»", "contractor": "BI Group", "contract": "№ 15 от 01.02.2026"}).json()
    assert a1["number"] == 1 and a1["total"] == pytest.approx(expected, abs=0.05)
    assert all(i["pct_before"] == 0 for i in a1["items"])
    assert client.post(f"/api/projects/{pid}/ks-acts", headers=h, json={"period_from": "2026-05-15", "period_to": "2026-06-30"}).status_code == 400
    assert client.post(f"/api/projects/{pid}/ks-acts", headers=h, json={"period_from": "2026-06-01", "period_to": "2026-06-30"}).status_code == 400  # nothing new

    # Progress on one task -> act 2 contains just the delta.
    t = next(t for t in leaves if 0 < t["percent_complete"] < 100)
    client.patch(f"/api/schedule/tasks/{t['id']}", headers=h, json={"percent_complete": 100})
    a2 = client.post(f"/api/projects/{pid}/ks-acts", headers=h, json={"period_from": "2026-06-01", "period_to": "2026-06-30"}).json()
    assert a2["number"] == 2 and len(a2["items"]) == 1
    it = a2["items"][0]
    assert it["pct_before"] == t["percent_complete"] and it["pct_now"] == 100
    assert a2["total"] == pytest.approx(t["cost"] * (100 - t["percent_complete"]) / 100, abs=0.01)
    assert a2["header"]["contractor"] == "BI Group"  # parties carried over

    ws = load_workbook(io.BytesIO(client.get(f"/api/ks-acts/{a2['id']}/ks2.xlsx", headers=h).content)).active
    cells = {str(c.value): c for row in ws.iter_rows() for c in row if c.value is not None}
    assert "АКТ № 2 О ПРИЁМКЕ ВЫПОЛНЕННЫХ РАБОТ" in cells and "BI Group" in cells
    vat_row = cells["НДС 12 %"].row
    assert ws.cell(row=vat_row, column=8).value == pytest.approx(a2["total"] * 0.12, abs=0.01)

    ws3 = load_workbook(io.BytesIO(client.get(f"/api/ks-acts/{a2['id']}/ks3.xlsx", headers=h).content)).active
    assert ws3["B11"].value == pytest.approx(a1["total"] + a2["total"], abs=0.05)  # since start
    assert ws3["D11"].value == pytest.approx(a2["total"], abs=0.01)  # period

    assert client.delete(f"/api/ks-acts/{a1['id']}", headers=h).status_code == 400  # only the last one
    assert client.delete(f"/api/ks-acts/{a2['id']}", headers=h).status_code == 204
    assert [a["number"] for a in client.get(f"/api/projects/{pid}/ks-acts", headers=h).json()] == [1]
