"""Issue register (замечания / предписания): workflow, assignee notifications,
comments, photo with EXIF GPS, filters, Excel/PDF export; org user admin."""

import io
from datetime import date, timedelta

import pytest


def _admin(client):
    r = client.post("/api/auth/register", json={"email": "boss@example.com", "password": "supersecret123",
                                                "full_name": "Начальник участка", "organization_name": "BI"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture()
def outbox(monkeypatch):
    from app import notify

    box = []
    monkeypatch.setattr(notify, "outbox", box)
    return box


def _worker(client, h, pid):
    r = client.post("/api/organization/users", headers=h, json={
        "email": "worker@example.com", "password": "worker12345", "full_name": "Прораб", "role": "geologist"})
    assert r.status_code == 201, r.text
    uid = r.json()["id"]
    client.post(f"/api/projects/{pid}/members", headers=h, json={"email": "worker@example.com", "role": "editor"})
    tok = client.post("/api/auth/login", json={"email": "worker@example.com", "password": "worker12345"}).json()["access_token"]
    return uid, {"Authorization": f"Bearer {tok}"}


def _jpeg_with_gps(lat: float, lon: float) -> bytes:
    from PIL import Image

    img = Image.new("RGB", (32, 32), (200, 50, 50))
    exif = Image.Exif()
    gps = {1: "N", 2: (int(lat), int((lat % 1) * 60), ((lat * 60) % 1) * 60), 3: "E",
           4: (int(lon), int((lon % 1) * 60), ((lon * 60) % 1) * 60)}
    exif[0x8825] = gps
    exif[306] = "2026:05:20 10:30:00"
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif)
    return buf.getvalue()


def test_issue_workflow(client, outbox):
    h = _admin(client)
    pid = client.post("/api/projects", json={"name": "ЖК"}, headers=h).json()["id"]
    worker_id, wh = _worker(client, h, pid)
    members = client.get(f"/api/projects/{pid}/members", headers=h).json()
    assert {m["name"] for m in members} == {"Начальник участка", "Прораб"}

    r = client.post(f"/api/projects/{pid}/issues", headers=h, json={
        "title": "Оголение арматуры", "kind": "prescription", "priority": "high", "assignee_id": worker_id,
        "due_date": str(date.today() - timedelta(days=1)), "longitude": 71.4, "latitude": 51.1, "height": 350.0,
    })
    assert r.status_code == 201, r.text
    issue = r.json()
    assert issue["number"] == 1 and issue["status"] == "open" and issue["overdue"] is True
    assert issue["assignee_name"] == "Прораб"
    assert outbox and "worker@example.com" in outbox[-1]["To"] and "Предписание №1" in outbox[-1]["Subject"]

    second = client.post(f"/api/projects/{pid}/issues", headers=h, json={"title": "Мусор на этаже"}).json()
    assert second["number"] == 2

    # Invalid jump open -> review
    assert client.patch(f"/api/issues/{issue['id']}", headers=wh, json={"status": "review"}).status_code == 400
    for st in ("in_progress", "review"):
        r = client.patch(f"/api/issues/{issue['id']}", headers=wh, json={"status": st})
        assert r.status_code == 200, r.text
    assert "boss@example.com" in outbox[-1]["To"]  # author notified, actor (worker) not

    r = client.post(f"/api/issues/{issue['id']}/comments", headers=wh, json={"text": "Устранено, прошу проверить"})
    photo = _jpeg_with_gps(51.123456, 71.654321)
    r = client.post(f"/api/issues/{second['id']}/attachments", headers=h, files={"file": ("site.jpg", photo, "image/jpeg")})
    det = r.json()
    assert det["latitude"] == pytest.approx(51.123456, abs=1e-4)  # location taken from the photo
    att = det["attachments"][0]
    assert att["taken_at"].startswith("2026-05-20")
    img = client.get(f"/api/issue-attachments/{att['id']}", headers=h)
    assert img.status_code == 200 and img.content == photo

    r = client.patch(f"/api/issues/{issue['id']}", headers=h, json={"status": "closed"})
    det = r.json()
    assert det["status"] == "closed" and det["closed_at"] and det["overdue"] is False
    kinds = [e["kind"] for e in det["events"]]
    assert kinds == ["created", "status", "status", "comment", "status"]

    assert [i["number"] for i in client.get(f"/api/projects/{pid}/issues", params={"status": "open"}, headers=h).json()] == [2]
    assert client.get(f"/api/projects/{pid}/issues", params={"assignee": "me"}, headers=wh).json()[0]["number"] == 1
    assert client.get(f"/api/projects/{pid}/issues", params={"q": "арматур"}, headers=h).json()[0]["number"] == 1

    x = client.get(f"/api/projects/{pid}/issues/export.xlsx", headers=h)
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A3"].value == "№" and {ws["C4"].value, ws["C5"].value} == {"Оголение арматуры", "Мусор на этаже"}
    p = client.get(f"/api/projects/{pid}/issues/export.pdf", headers=h)
    assert p.content[:4] == b"%PDF"

    # Viewer of another org cannot see
    other = client.post("/api/auth/register", json={"email": "x@example.com", "password": "supersecret123",
                                                    "full_name": "X", "organization_name": "Other"}).json()["access_token"]
    assert client.get(f"/api/issues/{issue['id']}", headers={"Authorization": f"Bearer {other}"}).status_code == 404


def test_org_admin_only(client):
    h = _admin(client)
    pid = client.post("/api/projects", json={"name": "ЖК"}, headers=h).json()["id"]
    _, wh = _worker(client, h, pid)
    assert client.get("/api/organization/users", headers=wh).status_code == 403
    assert len(client.get("/api/organization/users", headers=h).json()) == 2
    assert client.post("/api/organization/users", headers=h, json={"email": "worker@example.com", "password": "12345678"}).status_code == 400
