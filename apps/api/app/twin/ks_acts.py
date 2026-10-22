"""Drafts of the acts of completed works — КС-2 (акт о приёмке выполненных
работ) and КС-3 (справка о стоимости) — from the 4D schedule.

Amount per task = cost × (progress now − progress in the previous act);
tasks are matched between acts by external id / WBS / name, so a
re-imported schedule keeps the history. VAT (НДС РК) 12 % on top. Output:
Excel close to the unified forms, to be checked and posted in 1С (the
direct 1С exchange needs the customer's configuration).
"""

from __future__ import annotations

import io
from datetime import date

VAT = 0.12


def task_key(t) -> str:
    return t.external_id or t.wbs or t.name


def build_items(tasks: list, previous: dict[str, float]) -> list[dict]:
    items = []
    for t in sorted(tasks, key=lambda t: t.sort_order):
        if t.is_summary or not t.cost:
            continue
        key = task_key(t)
        before = float(previous.get(key, 0.0))
        now = float(t.percent_complete)
        if now <= before + 1e-9:
            continue
        items.append({
            "key": key, "name": t.name, "wbs": t.wbs, "cost": float(t.cost),
            "pct_before": round(before, 2), "pct_now": round(now, 2),
            "amount": round(float(t.cost) * (now - before) / 100.0, 2),
        })
    return items


def _money(ws, cell):
    ws[cell].number_format = "#,##0.00"


def ks2_xlsx(act, project_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "КС-2"
    h = act.header or {}
    thin = Side(style="thin", color="999999")
    ws["H1"] = "Унифицированная форма № КС-2 (черновик)"
    ws["H1"].font = Font(italic=True, size=9)
    rows = [("Заказчик", h.get("customer")), ("Подрядчик", h.get("contractor")), ("Стройка", project_name),
            ("Объект", h.get("object") or project_name), ("Договор", h.get("contract"))]
    for i, (k, v) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v or "")
    ws["A8"] = f"АКТ № {act.number} О ПРИЁМКЕ ВЫПОЛНЕННЫХ РАБОТ"
    ws["A8"].font = Font(bold=True, size=13)
    ws["A9"] = f"Отчётный период: с {act.period_from:%d.%m.%Y} по {act.period_to:%d.%m.%Y}"
    head = ["№ п/п", "Шифр / WBS", "Наименование работ", "Ед. изм.", "Выполнено, %", "Сметная стоимость, ₸",
            "Выполнено ранее, %", "Стоимость за период, ₸"]
    for c, v in enumerate(head, start=1):
        cell = ws.cell(row=11, column=c, value=v)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    r = 12
    for n, it in enumerate(act.items, start=1):
        vals = [n, it.get("wbs") or "", it["name"], "%", round(it["pct_now"] - it["pct_before"], 2), it["cost"], it["pct_before"], it["amount"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        _money(ws, f"F{r}")
        _money(ws, f"H{r}")
        r += 1
    total = round(sum(it["amount"] for it in act.items), 2)
    for label, val in (("Итого", total), (f"НДС {int(VAT * 100)} %", round(total * VAT, 2)), ("Всего с НДС", round(total * (1 + VAT), 2))):
        ws.cell(row=r, column=7, value=label).font = Font(bold=True)
        ws.cell(row=r, column=8, value=val).font = Font(bold=True)
        _money(ws, f"H{r}")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Сдал (подрядчик): ____________________")
    ws.cell(row=r, column=5, value="Принял (заказчик): ____________________")
    for col, w in zip("ABCDEFGH", (6, 14, 48, 8, 12, 18, 14, 20)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ks3_xlsx(act, project_name: str, since_start: float, since_year: float) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "КС-3"
    h = act.header or {}
    ws["D1"] = "Унифицированная форма № КС-3 (черновик)"
    ws["D1"].font = Font(italic=True, size=9)
    for i, (k, v) in enumerate((("Заказчик", h.get("customer")), ("Подрядчик", h.get("contractor")), ("Стройка", project_name),
                                ("Договор", h.get("contract"))), start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v or "")
    ws["A7"] = f"СПРАВКА № {act.number} О СТОИМОСТИ ВЫПОЛНЕННЫХ РАБОТ И ЗАТРАТ"
    ws["A7"].font = Font(bold=True, size=13)
    ws["A8"] = f"Отчётный период: с {act.period_from:%d.%m.%Y} по {act.period_to:%d.%m.%Y}"
    head = ["", "С начала строительства, ₸", "С начала года, ₸", "За отчётный период, ₸"]
    for c, v in enumerate(head, start=1):
        ws.cell(row=10, column=c, value=v).font = Font(bold=True)
    period = float(act.total)
    for r, (label, k) in enumerate((("Стоимость работ", 1.0), (f"НДС {int(VAT * 100)} %", VAT), ("Всего с НДС", 1 + VAT)), start=11):
        ws.cell(row=r, column=1, value=label)
        for c, v in enumerate((since_start, since_year, period), start=2):
            ws.cell(row=r, column=c, value=round(v * k, 2)).number_format = "#,##0.00"
    for col, w in zip("ABCD", (28, 26, 22, 24)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def year_start(d: date) -> date:
    return date(d.year, 1, 1)
