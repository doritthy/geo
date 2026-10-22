"""Work schedules for 4D plan-fact: import (MS Project XML, Excel/CSV) and
earned-value style progress metrics.

MS Project's binary .mpp is not parsed (it needs the Java MPXJ library);
MS Project / Primavera export to XML ("Сохранить как → XML") is supported.
"""

from __future__ import annotations

import csv
import io
import math
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta


@dataclass
class ParsedTask:
    external_id: str | None
    name: str
    start: date
    finish: date
    percent: float = 0.0
    cost: float | None = None
    wbs: str | None = None
    outline_level: int = 1
    is_summary: bool = False
    children: list["ParsedTask"] = field(default_factory=list)


def _date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text[:19] if "T" in text else text, fmt).date()
        except ValueError:
            continue
    return None


def _num(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace("%", "").replace(" ", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def parse_msproject_xml(data: bytes) -> list[ParsedTask]:
    root = ET.fromstring(data)
    ns = {"p": root.tag.split("}")[0].strip("{")} if root.tag.startswith("{") else {}

    def find(el, tag):
        return el.find(f"p:{tag}", ns) if ns else el.find(tag)

    def text(el, tag):
        node = find(el, tag)
        return node.text if node is not None else None

    tasks_el = find(root, "Tasks")
    if tasks_el is None:
        raise ValueError("В XML нет раздела Tasks — это не файл MS Project XML")
    out: list[ParsedTask] = []
    for t in (tasks_el.findall("p:Task", ns) if ns else tasks_el.findall("Task")):
        if text(t, "UID") == "0" or text(t, "IsNull") == "1":
            continue  # project summary row / blank
        name = (text(t, "Name") or "").strip()
        start, finish = _date(text(t, "Start")), _date(text(t, "Finish"))
        if not name or not start or not finish:
            continue
        out.append(ParsedTask(
            external_id=text(t, "UID"), name=name, start=start, finish=finish,
            percent=_num(text(t, "PercentComplete")) or 0.0, cost=_num(text(t, "Cost")),
            wbs=text(t, "WBS"), outline_level=int(text(t, "OutlineLevel") or 1),
            is_summary=text(t, "Summary") == "1",
        ))
    return out


_COLUMNS = {
    "id": ("id", "№", "ид", "uid", "номер", "n"),
    "wbs": ("wbs", "сдр", "код", "шифр"),
    "name": ("наименование", "название", "задача", "работа", "name", "task name", "task"),
    "start": ("начало", "дата начала", "start", "start date"),
    "finish": ("окончание", "конец", "дата окончания", "finish", "end", "finish date"),
    "percent": ("% выполнения", "%", "выполнено", "процент выполнения", "percent complete", "% complete", "progress"),
    "cost": ("стоимость", "сумма", "cost", "бюджет"),
    "level": ("уровень", "outline level", "level"),
}


def _header_map(row: list) -> dict[str, int]:
    found = {}
    for i, cell in enumerate(row):
        key = re.sub(r"\s+", " ", str(cell or "").strip().lower())
        for col, names in _COLUMNS.items():
            if col not in found and key in names:
                found[col] = i
    return found


def parse_table(rows: list[list]) -> list[ParsedTask]:
    header_idx, cols = None, {}
    for i, row in enumerate(rows[:20]):
        cols = _header_map(row)
        if {"name", "start", "finish"} <= cols.keys():
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Не найдены колонки «Наименование», «Начало», «Окончание»")
    out = []
    for row in rows[header_idx + 1:]:
        def get(col):
            i = cols.get(col)
            return row[i] if i is not None and i < len(row) else None

        name = str(get("name") or "").strip()
        start, finish = _date(get("start")), _date(get("finish"))
        if not name or not start or not finish:
            continue
        pct = _num(get("percent")) or 0.0
        if 0 < pct <= 1 and isinstance(get("percent"), float):
            pct *= 100  # Excel stores 45% as 0.45
        level = int(_num(get("level")) or 1)
        out.append(ParsedTask(external_id=str(get("id")) if get("id") not in (None, "") else None, name=name,
                              start=start, finish=finish, percent=pct, cost=_num(get("cost")),
                              wbs=str(get("wbs")) if get("wbs") else None, outline_level=level))
    return out


def parse_schedule(filename: str, data: bytes) -> list[ParsedTask]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "mpp":
        raise ValueError("Формат .mpp не поддерживается: сохраните график в MS Project как XML (Файл → Сохранить как → XML)")
    if ext == "xml":
        tasks = parse_msproject_xml(data)
    elif ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).active
        tasks = parse_table([list(r) for r in ws.iter_rows(values_only=True)])
    elif ext in ("csv", "txt"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t")
        tasks = parse_table(list(csv.reader(io.StringIO(text), dialect)))
    else:
        raise ValueError("Поддерживаются MS Project XML, Excel (.xlsx) и CSV")
    if not tasks:
        raise ValueError("В файле не найдено задач с датами")
    # Summary flag from outline levels when the source does not provide it.
    for i, t in enumerate(tasks):
        if not t.is_summary and i + 1 < len(tasks) and tasks[i + 1].outline_level > t.outline_level:
            t.is_summary = True
    return tasks


# ---------- metrics ----------


def planned_percent(start: date, finish: date, on: date) -> float:
    if on < start:
        return 0.0
    if on >= finish:
        return 100.0
    total = (finish - start).days + 1
    return 100.0 * ((on - start).days + 1) / total


def task_status(start: date, finish: date, actual: float, on: date) -> str:
    """not_started | in_progress | done | late (behind plan) | future."""
    if actual >= 100:
        return "done"
    plan = planned_percent(start, finish, on)
    if on > finish or plan - actual > 10:
        return "late"
    if actual > 0:
        return "in_progress"
    return "future" if on < start else "not_started"


def metrics(tasks: list, on: date) -> dict:
    """Earned-value style metrics over leaf tasks. Weight = cost when every
    leaf has one, otherwise duration in days."""
    leaves = [t for t in tasks if not t.is_summary]
    if not leaves:
        return {"status_date": on.isoformat()}
    use_cost = all((t.cost or 0) > 0 for t in leaves)
    total_w = pv = ev = 0.0
    for t in leaves:
        w = t.cost if use_cost else float((t.finish_date - t.start_date).days + 1)
        total_w += w
        pv += w * planned_percent(t.start_date, t.finish_date, on) / 100
        ev += w * t.percent_complete / 100
    start = min(t.start_date for t in leaves)
    finish = max(t.finish_date for t in leaves)
    spi = ev / pv if pv > 0 else None
    # Forecast: remaining work at the pace achieved so far, counted from the
    # status date (unlike duration/SPI this stays valid after the planned
    # finish, when PV is saturated, and never lands in the past).
    forecast = None
    if ev >= total_w:
        forecast = on
    elif ev > 0 and on > start:
        rate = ev / (on - start).days
        forecast = on + timedelta(days=math.ceil((total_w - ev) / rate))
    return {
        "status_date": on.isoformat(),
        "weight": "cost" if use_cost else "duration",
        "planned_percent": 100 * pv / total_w if total_w else 0,
        "actual_percent": 100 * ev / total_w if total_w else 0,
        "pv": pv, "ev": ev, "spi": spi,
        "start": start.isoformat(), "finish": finish.isoformat(),
        "forecast_finish": forecast.isoformat() if forecast else None,
        "delay_days": (forecast - finish).days if forecast else None,
        "late_tasks": sum(1 for t in leaves if task_status(t.start_date, t.finish_date, t.percent_complete, on) == "late"),
        "tasks": len(leaves),
    }
