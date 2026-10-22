"""Замечания и предписания: issue register with workflow
(Открыто → В работе → На проверке → Закрыто), assignee, due date, priority,
photo evidence (EXIF GPS), comments/history, Excel/PDF export and email
notifications."""

import io
import uuid
from datetime import date, datetime, timezone
from urllib.parse import quote

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db import get_db
from app.deps import _require_project_role, get_project_scoped_or_404, project_editor, project_viewer
from app.models import Issue, IssueAttachment, IssueEvent, Project, ProjectMember, SceneAsset, User
from app.notify import send_email
from app.security import get_current_user
from app.storage import open_stream, upload_fileobj

router = APIRouter(tags=["issues"])

STATUSES = ("open", "in_progress", "review", "closed")
STATUS_LABELS = {"open": "Открыто", "in_progress": "В работе", "review": "На проверке", "closed": "Закрыто"}
PRIORITY_LABELS = {"low": "Низкий", "normal": "Обычный", "high": "Высокий", "critical": "Критический"}
KIND_LABELS = {"remark": "Замечание", "prescription": "Предписание"}
# Allowed workflow moves.
TRANSITIONS = {
    "open": {"in_progress", "closed"},
    "in_progress": {"review", "open"},
    "review": {"closed", "in_progress"},
    "closed": {"open"},
}


class IssueCreate(BaseModel):
    title: str = Field(min_length=1, max_length=300)
    description: str | None = None
    kind: str = Field(default="remark", pattern="^(remark|prescription)$")
    priority: str = Field(default="normal", pattern="^(low|normal|high|critical)$")
    assignee_id: uuid.UUID | None = None
    due_date: date | None = None
    longitude: float | None = Field(default=None, ge=-180, le=180)
    latitude: float | None = Field(default=None, ge=-90, le=90)
    height: float | None = None
    asset_id: uuid.UUID | None = None
    element_global_id: str | None = None


class IssueUpdate(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=300)
    description: str | None = None
    kind: str | None = Field(default=None, pattern="^(remark|prescription)$")
    priority: str | None = Field(default=None, pattern="^(low|normal|high|critical)$")
    assignee_id: uuid.UUID | None = None
    due_date: date | None = None
    status: str | None = Field(default=None, pattern="^(open|in_progress|review|closed)$")
    longitude: float | None = None
    latitude: float | None = None
    height: float | None = None


class CommentCreate(BaseModel):
    text: str = Field(min_length=1, max_length=5000)


class AttachmentOut(BaseModel):
    id: uuid.UUID
    filename: str
    content_type: str | None
    size_bytes: int
    longitude: float | None
    latitude: float | None
    taken_at: datetime | None
    created_at: datetime


class EventOut(BaseModel):
    id: int
    kind: str
    data: dict
    user_name: str | None
    created_at: datetime


class IssueOut(BaseModel):
    id: uuid.UUID
    number: int
    kind: str
    title: str
    description: str | None
    status: str
    priority: str
    assignee_id: uuid.UUID | None
    assignee_name: str | None
    due_date: date | None
    overdue: bool
    longitude: float | None
    latitude: float | None
    height: float | None
    asset_id: uuid.UUID | None
    element_global_id: str | None
    created_by: uuid.UUID | None
    author_name: str | None
    created_at: datetime
    updated_at: datetime
    closed_at: datetime | None
    attachments_count: int


class IssueDetail(IssueOut):
    attachments: list[AttachmentOut]
    events: list[EventOut]


class MemberOut(BaseModel):
    id: uuid.UUID
    name: str
    email: str
    role: str


# ---------- helpers ----------


def _names(db: Session, ids: set) -> dict:
    ids = {i for i in ids if i}
    return {u.id: (u.full_name or u.email) for u in db.query(User).filter(User.id.in_(ids)).all()} if ids else {}


def _out(db: Session, issues: list[Issue]) -> list[IssueOut]:
    names = _names(db, {i.assignee_id for i in issues} | {i.created_by for i in issues})
    today = date.today()
    return [
        IssueOut(
            id=i.id, number=i.number, kind=i.kind, title=i.title, description=i.description, status=i.status,
            priority=i.priority, assignee_id=i.assignee_id, assignee_name=names.get(i.assignee_id),
            due_date=i.due_date, overdue=bool(i.due_date and i.due_date < today and i.status != "closed"),
            longitude=i.longitude, latitude=i.latitude, height=i.height, asset_id=i.asset_id,
            element_global_id=i.element_global_id, created_by=i.created_by, author_name=names.get(i.created_by),
            created_at=i.created_at, updated_at=i.updated_at, closed_at=i.closed_at,
            attachments_count=len(i.attachments),
        )
        for i in issues
    ]


def _detail(db: Session, issue: Issue) -> IssueDetail:
    events = db.query(IssueEvent).filter(IssueEvent.issue_id == issue.id).order_by(IssueEvent.id).all()
    names = _names(db, {e.user_id for e in events})
    base = _out(db, [issue])[0]
    return IssueDetail(
        **base.model_dump(),
        attachments=[AttachmentOut.model_validate(a, from_attributes=True) for a in issue.attachments],
        events=[EventOut(id=e.id, kind=e.kind, data=e.data, user_name=names.get(e.user_id), created_at=e.created_at)
                for e in events],
    )


def _project_users(db: Session, project: Project) -> list[User]:
    member_ids = db.query(ProjectMember.user_id).filter(ProjectMember.project_id == project.id).subquery()
    return (
        db.query(User)
        .filter(User.organization_id == project.organization_id, or_(User.id.in_(member_ids), User.role == "admin"))
        .order_by(User.full_name, User.email)
        .all()
    )


def _check_assignee(db: Session, project: Project, user_id: uuid.UUID | None) -> None:
    if user_id is None:
        return
    if user_id not in {u.id for u in _project_users(db, project)}:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Исполнитель не участник проекта")


def _issue_for(issue_id: uuid.UUID, db: Session, user: User, editor: bool = False) -> Issue:
    issue: Issue = get_project_scoped_or_404(Issue, issue_id, db, user)
    if editor:
        _require_project_role(issue.project_id, "editor", db, user)
    return issue


def _notify(bg: BackgroundTasks, db: Session, issue: Issue, actor: User, subject: str, body: str) -> None:
    ids = {issue.assignee_id, issue.created_by} - {actor.id, None}
    emails = [u.email for u in db.query(User).filter(User.id.in_(ids)).all()] if ids else []
    if not emails:
        return
    link = f"{get_settings().PUBLIC_WEB_URL.rstrip('/')}/twin/{issue.project_id}?issue={issue.id}"
    bg.add_task(send_email, emails, f"[{KIND_LABELS[issue.kind]} №{issue.number}] {subject}",
                f"{body}\n\n«{issue.title}»\nСтатус: {STATUS_LABELS[issue.status]}\n"
                f"Срок: {issue.due_date.strftime('%d.%m.%Y') if issue.due_date else '—'}\n\n{link}\n")


def _exif_gps(data: bytes) -> tuple[float | None, float | None, datetime | None]:
    try:
        from PIL import Image

        img = Image.open(io.BytesIO(data))
        exif = img.getexif()
        taken = None
        dt = exif.get(306) or exif.get_ifd(0x8769).get(36867)
        if dt:
            taken = datetime.strptime(str(dt), "%Y:%m:%d %H:%M:%S")
        gps = exif.get_ifd(0x8825)
        if not gps or 2 not in gps or 4 not in gps:
            return None, None, taken

        def deg(v, ref):
            d = float(v[0]) + float(v[1]) / 60 + float(v[2]) / 3600
            return -d if ref in ("S", "W") else d

        return deg(gps[4], gps.get(3, "E")), deg(gps[2], gps.get(1, "N")), taken
    except Exception:  # noqa: BLE001 -- not an image / no EXIF
        return None, None, None


# ---------- endpoints ----------


@router.get("/api/projects/{project_id}/members", response_model=list[MemberOut])
def members(project: Project = Depends(project_viewer), db: Session = Depends(get_db)):
    roles = {m.user_id: m.role for m in db.query(ProjectMember).filter(ProjectMember.project_id == project.id).all()}
    return [MemberOut(id=u.id, name=u.full_name or u.email, email=u.email, role=roles.get(u.id, "admin"))
            for u in _project_users(db, project)]


def _filtered(db: Session, project: Project, status_: str | None, assignee: str | None, kind: str | None,
              overdue: bool, q: str | None, user: User):
    query = db.query(Issue).filter(Issue.project_id == project.id)
    if status_:
        query = query.filter(Issue.status.in_(status_.split(",")))
    if kind:
        query = query.filter(Issue.kind == kind)
    if assignee == "me":
        query = query.filter(Issue.assignee_id == user.id)
    elif assignee:
        query = query.filter(Issue.assignee_id == uuid.UUID(assignee))
    if overdue:
        query = query.filter(Issue.due_date < date.today(), Issue.status != "closed")
    if q:
        like = f"%{q}%"
        cond = [Issue.title.ilike(like), Issue.description.ilike(like)]
        if q.isdigit():
            cond.append(Issue.number == int(q))
        query = query.filter(or_(*cond))
    return query.order_by(Issue.number.desc()).all()


@router.get("/api/projects/{project_id}/issues", response_model=list[IssueOut])
def list_issues(
    status_: str | None = Query(None, alias="status"),
    assignee: str | None = None,
    kind: str | None = None,
    overdue: bool = False,
    q: str | None = None,
    project: Project = Depends(project_viewer),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    return _out(db, _filtered(db, project, status_, assignee, kind, overdue, q, user))


@router.post("/api/projects/{project_id}/issues", response_model=IssueDetail, status_code=status.HTTP_201_CREATED)
def create_issue(payload: IssueCreate, background_tasks: BackgroundTasks, project: Project = Depends(project_editor),
                 db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    _check_assignee(db, project, payload.assignee_id)
    if payload.asset_id:
        a = db.get(SceneAsset, payload.asset_id)
        if a is None or a.project_id != project.id:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Asset not found")
    # Serialize numbering per project.
    db.query(Project).filter(Project.id == project.id).with_for_update().one()
    number = (db.query(func.max(Issue.number)).filter(Issue.project_id == project.id).scalar() or 0) + 1
    issue = Issue(project_id=project.id, number=number, created_by=user.id, **payload.model_dump())
    db.add(issue)
    db.flush()
    db.add(IssueEvent(issue_id=issue.id, user_id=user.id, kind="created", data={}))
    db.commit()
    db.refresh(issue)
    if issue.assignee_id:
        _notify(background_tasks, db, issue, user, "Вам назначено", f"{user.full_name or user.email} назначил(а) вам {KIND_LABELS[issue.kind].lower()}.")
    return _detail(db, issue)


@router.get("/api/issues/{issue_id}", response_model=IssueDetail)
def get_issue(issue_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return _detail(db, _issue_for(issue_id, db, user))


@router.patch("/api/issues/{issue_id}", response_model=IssueDetail)
def update_issue(issue_id: uuid.UUID, payload: IssueUpdate, background_tasks: BackgroundTasks,
                 db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    issue = _issue_for(issue_id, db, user, editor=True)
    data = payload.model_dump(exclude_unset=True)
    project = db.get(Project, issue.project_id)

    if "status" in data and data["status"] != issue.status:
        new = data.pop("status")
        if new not in TRANSITIONS[issue.status]:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Переход «{STATUS_LABELS[issue.status]}» → «{STATUS_LABELS[new]}» недопустим",
            )
        db.add(IssueEvent(issue_id=issue.id, user_id=user.id, kind="status", data={"from": issue.status, "to": new}))
        issue.status = new
        issue.closed_at = datetime.now(timezone.utc) if new == "closed" else None
        _notify(background_tasks, db, issue, user, f"Статус: {STATUS_LABELS[new]}",
                f"{user.full_name or user.email} изменил(а) статус на «{STATUS_LABELS[new]}».")
    else:
        data.pop("status", None)

    if "assignee_id" in data and data["assignee_id"] != issue.assignee_id:
        _check_assignee(db, project, data["assignee_id"])
        issue.assignee_id = data.pop("assignee_id")
        db.add(IssueEvent(issue_id=issue.id, user_id=user.id, kind="assigned",
                          data={"assignee_id": str(issue.assignee_id) if issue.assignee_id else None}))
        if issue.assignee_id:
            _notify(background_tasks, db, issue, user, "Вам назначено",
                    f"{user.full_name or user.email} назначил(а) вам {KIND_LABELS[issue.kind].lower()}.")
    else:
        data.pop("assignee_id", None)

    changed = {}
    for key, value in data.items():
        if getattr(issue, key) != value:
            changed[key] = str(value) if value is not None else None
            setattr(issue, key, value)
    if changed:
        db.add(IssueEvent(issue_id=issue.id, user_id=user.id, kind="updated", data=changed))
    db.commit()
    db.refresh(issue)
    return _detail(db, issue)


@router.delete("/api/issues/{issue_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_issue(issue_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    issue = _issue_for(issue_id, db, user)
    _require_project_role(issue.project_id, "owner", db, user)
    db.delete(issue)
    db.commit()


@router.post("/api/issues/{issue_id}/comments", response_model=IssueDetail, status_code=status.HTTP_201_CREATED)
def add_comment(issue_id: uuid.UUID, payload: CommentCreate, background_tasks: BackgroundTasks,
                db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    issue = _issue_for(issue_id, db, user)  # viewers may comment
    db.add(IssueEvent(issue_id=issue.id, user_id=user.id, kind="comment", data={"text": payload.text}))
    db.commit()
    _notify(background_tasks, db, issue, user, "Новый комментарий", f"{user.full_name or user.email}: {payload.text}")
    return _detail(db, issue)


@router.post("/api/issues/{issue_id}/attachments", response_model=IssueDetail, status_code=status.HTTP_201_CREATED)
def add_attachment(issue_id: uuid.UUID, file: UploadFile = File(...), use_location: bool = Form(True),
                   db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    issue = _issue_for(issue_id, db, user, editor=True)
    data = file.file.read()
    lon, lat, taken = _exif_gps(data) if (file.content_type or "").startswith("image/") else (None, None, None)
    att_id = uuid.uuid4()
    key = f"projects/{issue.project_id}/issues/{issue.id}/{att_id}/{file.filename or 'file'}"
    upload_fileobj(key, io.BytesIO(data), file.content_type or "application/octet-stream")
    db.add(IssueAttachment(id=att_id, issue_id=issue.id, filename=file.filename or "file", storage_key=key,
                           content_type=file.content_type, size_bytes=len(data), longitude=lon, latitude=lat,
                           taken_at=taken, uploaded_by=user.id))
    if use_location and lon is not None and issue.longitude is None:
        issue.longitude, issue.latitude = lon, lat
    db.add(IssueEvent(issue_id=issue.id, user_id=user.id, kind="attachment",
                      data={"filename": file.filename, "gps": lon is not None}))
    db.commit()
    db.refresh(issue)
    return _detail(db, issue)


@router.get("/api/issue-attachments/{attachment_id}")
def download_attachment(attachment_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    att = db.get(IssueAttachment, attachment_id)
    if att is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")
    _issue_for(att.issue_id, db, user)
    body, length = open_stream(att.storage_key)
    headers = {"Content-Disposition": f"inline; filename*=UTF-8''{quote(att.filename)}", "Cache-Control": "private, max-age=3600"}
    if length is not None:
        headers["Content-Length"] = str(length)
    return StreamingResponse(body.iter_chunks(256 * 1024), media_type=att.content_type or "application/octet-stream", headers=headers)


# ---------- export ----------


def _rows(db: Session, issues: list[Issue]) -> list[list]:
    out = _out(db, issues)
    return [
        [i.number, KIND_LABELS[i.kind], i.title, i.description or "", STATUS_LABELS[i.status], PRIORITY_LABELS[i.priority],
         i.assignee_name or "", i.due_date.strftime("%d.%m.%Y") if i.due_date else "", "да" if i.overdue else "",
         i.author_name or "", i.created_at.strftime("%d.%m.%Y"), i.closed_at.strftime("%d.%m.%Y") if i.closed_at else "",
         i.element_global_id or "",
         f"{i.latitude:.6f}, {i.longitude:.6f}" if i.latitude is not None and i.longitude is not None else ""]
        for i in out
    ]


HEADER = ["№", "Тип", "Наименование", "Описание", "Статус", "Приоритет", "Исполнитель", "Срок", "Просрочено",
          "Автор", "Создано", "Закрыто", "Элемент (GlobalId)", "Координаты"]


@router.get("/api/projects/{project_id}/issues/export.xlsx")
def export_xlsx(status_: str | None = Query(None, alias="status"), assignee: str | None = None, kind: str | None = None,
                overdue: bool = False, q: str | None = None, project: Project = Depends(project_viewer),
                db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр"
    ws.append([f"Реестр замечаний и предписаний — {project.name}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Сформирован {datetime.now().strftime('%d.%m.%Y %H:%M')}"])
    ws.append(HEADER)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EA")
    for row in _rows(db, _filtered(db, project, status_, assignee, kind, overdue, q, user)):
        ws.append(row)
    widths = [6, 13, 40, 50, 13, 12, 22, 11, 11, 22, 11, 11, 24, 24]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    name = f"Предписания_{date.today():%Y-%m-%d}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


@router.get("/api/projects/{project_id}/issues/export.pdf")
def export_pdf(status_: str | None = Query(None, alias="status"), assignee: str | None = None, kind: str | None = None,
               overdue: bool = False, q: str | None = None, project: Project = Depends(project_viewer),
               db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    from app.twin.report import _font

    font = _font()
    st = ParagraphStyle("c", fontName=font, fontSize=7.5, leading=9)
    h1 = ParagraphStyle("h", fontName=font, fontSize=13, leading=16, spaceAfter=6)
    rows = _rows(db, _filtered(db, project, status_, assignee, kind, overdue, q, user))
    cols = [0, 1, 2, 4, 5, 6, 7, 9, 10]
    data = [[Paragraph(HEADER[c], st) for c in cols]] + [[Paragraph(str(r[c]), st) for c in cols] for r in rows]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    t = Table(data, colWidths=[10 * mm, 22 * mm, 80 * mm, 22 * mm, 20 * mm, 40 * mm, 18 * mm, 40 * mm, 18 * mm], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    doc.build([Paragraph(f"Реестр замечаний и предписаний — {project.name}", h1),
               Paragraph(f"Сформирован {datetime.now().strftime('%d.%m.%Y %H:%M')}. Всего: {len(rows)}", st), t])
    name = f"Предписания_{date.today():%Y-%m-%d}.pdf"
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})
